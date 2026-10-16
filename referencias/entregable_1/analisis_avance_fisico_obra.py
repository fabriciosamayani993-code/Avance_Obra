# ============================================================
# ANALISIS DEL AVANCE FISICO DE UNA OBRA MEDIANTE PYTHON
# Caso 2 - Entregable Python + IA
# ============================================================

# Importar Pandas para leer, organizar y analizar datos tabulares
import pandas as pd

# Importar NumPy para realizar operaciones numéricas
import numpy as np

# Importar Matplotlib para crear gráficos
import matplotlib.pyplot as plt

# Importar FuncAnimation para crear una animación
from matplotlib.animation import FuncAnimation

# Importar Plotly para crear gráficos interactivos
import plotly.graph_objects as go

# Importar OpenPyXL para trabajar con Excel
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image

# Importar os para trabajar con rutas de archivos
import os


# ============================================================
# 1. RUTAS DEL PROYECTO
# ============================================================

# Obtener la carpeta donde se encuentra este archivo .py
BASE = os.path.dirname(os.path.abspath(__file__))

# Construir la ruta hacia el archivo CSV
ruta_csv = os.path.join(BASE, "base_datos_avance_obra.csv")

# Definir la ruta del reporte Excel
ruta_excel = os.path.join(BASE, "reporte_avance_obra.xlsx")

# Definir las rutas de los gráficos
ruta_curva_s = os.path.join(BASE, "curva_s.png")
ruta_produccion = os.path.join(BASE, "produccion_semanal.png")
ruta_especialidad = os.path.join(BASE, "avance_especialidad.png")
ruta_partidas_encargado = os.path.join(BASE, "partidas_encargado.png")
ruta_avance_encargado = os.path.join(BASE, "avance_encargado.png")


# ============================================================
# 2. LECTURA Y EXPLORACION DE DATOS
# ============================================================

# Leer el archivo CSV y almacenarlo en un DataFrame
df = pd.read_csv(ruta_csv)

# Mostrar información básica de la base de datos
print("===== BASE DE DATOS CARGADA =====")
print("Dimensiones:", df.shape)
print("Columnas:", list(df.columns))
print("\nPrimeras filas:")
print(df.head())


# ============================================================
# 3. VALIDACION DE LA BASE DE DATOS
# ============================================================

# Contar partidas, semanas, especialidades y encargados únicos
cantidad_partidas = df["codigo_partida"].nunique()
cantidad_semanas = df["semana"].nunique()
cantidad_especialidades = df["especialidad"].nunique()
cantidad_encargados = df["encargado"].nunique()

# Obtener una sola fila por partida para evitar repetir los pesos
pesos_unicos = df.drop_duplicates(subset="codigo_partida")

# Sumar los pesos de las partidas
peso_total = pesos_unicos["peso_partida"].sum()

# Mostrar los resultados
print("\n===== VALIDACION =====")
print("Cantidad de partidas:", cantidad_partidas)
print("Cantidad de semanas:", cantidad_semanas)
print("Cantidad de especialidades:", cantidad_especialidades)
print("Cantidad de encargados:", cantidad_encargados)
print(f"Suma de pesos: {peso_total:.2f} %")


# ============================================================
# 4. FUNCIONES DEL ANALISIS
# ============================================================

# Crear una función para calcular el avance semanal
def calcular_avance_semanal(datos):
    # Agrupar por semana y sumar los avances planificado y real
    resumen = datos.groupby("semana")[[
        "avance_plan_semanal",
        "avance_real_semanal"
    ]].sum().reset_index()

    # Redondear los resultados
    resumen[["avance_plan_semanal", "avance_real_semanal"]] = resumen[[
        "avance_plan_semanal",
        "avance_real_semanal"
    ]].round(2)

    # Devolver la tabla calculada
    return resumen


# Crear una función para calcular el avance acumulado
def calcular_avance_acumulado(datos_semanales):
    # Crear una copia para no modificar la tabla original
    resultado = datos_semanales.copy()

    # Calcular las sumas acumuladas
    resultado["avance_plan_acumulado"] = resultado["avance_plan_semanal"].cumsum()
    resultado["avance_real_acumulado"] = resultado["avance_real_semanal"].cumsum()

    # Redondear los valores
    resultado[["avance_plan_acumulado", "avance_real_acumulado"]] = resultado[[
        "avance_plan_acumulado",
        "avance_real_acumulado"
    ]].round(2)

    # Devolver la tabla
    return resultado


# Crear una función para calcular la diferencia y clasificar el estado
def calcular_diferencia_avance(datos_acumulados):
    # Crear una copia
    resultado = datos_acumulados.copy()

    # Calcular la diferencia: Real - Planificado
    resultado["diferencia"] = (
        resultado["avance_real_acumulado"] - resultado["avance_plan_acumulado"]
    ).round(2)

    # Evitar que aparezcan valores como -0.00
    resultado["diferencia"] = resultado["diferencia"].where(
        np.abs(resultado["diferencia"]) >= 0.005,
        0
    )

    # Clasificar el estado de la obra
    resultado["estado"] = np.where(
        resultado["diferencia"] > 0,
        "ADELANTADA",
        np.where(
            resultado["diferencia"] < 0,
            "RETRASADA",
            "SEGUN LO PREVISTO"
        )
    )

    # Devolver la tabla
    return resultado


# Crear una función para analizar el avance por especialidad
def analizar_avance_especialidad(datos):
    # Agrupar por especialidad y sumar los avances
    resumen = datos.groupby("especialidad")[[
        "avance_plan_semanal",
        "avance_real_semanal"
    ]].sum()

    # Calcular la diferencia
    resumen["diferencia"] = (
        resumen["avance_real_semanal"] - resumen["avance_plan_semanal"]
    )

    # Redondear y convertir el índice en columna
    resumen = resumen.round(2).reset_index()

    # Evitar valores -0.00
    resumen["diferencia"] = resumen["diferencia"].where(
        np.abs(resumen["diferencia"]) >= 0.005,
        0
    )

    # Devolver la tabla
    return resumen


# Crear una función para contar partidas por encargado
def contar_partidas_encargado(datos):
    # Conservar una sola fila por partida
    partidas_unicas = datos.drop_duplicates(subset="codigo_partida")

    # Agrupar por encargado y contar las partidas
    resumen = (
        partidas_unicas
        .groupby("encargado")["codigo_partida"]
        .count()
        .reset_index(name="cantidad_partidas")
    )

    # Devolver la tabla
    return resumen


# Crear una función para analizar el avance por encargado
def analizar_avance_encargado(datos):
    # Agrupar por encargado y sumar los avances
    resumen = datos.groupby("encargado")[[
        "avance_plan_semanal",
        "avance_real_semanal"
    ]].sum()

    # Calcular la diferencia
    resumen["diferencia"] = (
        resumen["avance_real_semanal"] - resumen["avance_plan_semanal"]
    )

    # Redondear y convertir el índice en columna
    resumen = resumen.round(2).reset_index()

    # Evitar valores -0.00
    resumen["diferencia"] = resumen["diferencia"].where(
        np.abs(resumen["diferencia"]) >= 0.005,
        0
    )

    # Devolver la tabla
    return resumen


# ============================================================
# 5. APLICACION DE LAS FUNCIONES
# ============================================================

# Calcular el avance semanal
avance_semanal = calcular_avance_semanal(df)

# Calcular los acumulados
avance_acumulado = calcular_avance_acumulado(avance_semanal)

# Calcular diferencias y estados
resumen_semanal = calcular_diferencia_avance(avance_acumulado)

# Calcular el avance por especialidad
avance_especialidad = analizar_avance_especialidad(df)

# Contar las partidas por encargado
partidas_encargado = contar_partidas_encargado(df)

# Calcular el avance por encargado
avance_encargado = analizar_avance_encargado(df)

# Mostrar los resultados principales
print("\n===== RESUMEN SEMANAL =====")
print(resumen_semanal.to_string(index=False))

print("\n===== AVANCE POR ESPECIALIDAD =====")
print(avance_especialidad.to_string(index=False))

print("\n===== AVANCE POR ENCARGADO =====")
print(avance_encargado.to_string(index=False))


# ============================================================
# 6. CURVA S INTERACTIVA CON PLOTLY
# ============================================================

# Crear la figura interactiva
fig = go.Figure()

# Agregar la curva planificada
fig.add_trace(go.Scatter(
    x=resumen_semanal["semana"],
    y=resumen_semanal["avance_plan_acumulado"],
    mode="lines+markers",
    name="Planificado"
))

# Agregar la curva real
fig.add_trace(go.Scatter(
    x=resumen_semanal["semana"],
    y=resumen_semanal["avance_real_acumulado"],
    mode="lines+markers",
    name="Real"
))

# Personalizar el gráfico
fig.update_layout(
    title="Curva S - Avance planificado vs. real",
    xaxis_title="Semana",
    yaxis_title="Avance acumulado (%)",
    hovermode="x unified"
)

# Mostrar el gráfico interactivo
fig.show()


# ============================================================
# 7. GRAFICOS ESTATICOS CON MATPLOTLIB
# ============================================================

# Crear y guardar la Curva S
plt.figure(figsize=(8, 4.5))
plt.plot(
    resumen_semanal["semana"],
    resumen_semanal["avance_plan_acumulado"],
    marker="o",
    linestyle="--",
    linewidth=2,
    label="Planificado"
)
plt.plot(
    resumen_semanal["semana"],
    resumen_semanal["avance_real_acumulado"],
    marker="s",
    linewidth=2,
    label="Real"
)
plt.title("Curva S - Avance planificado vs. real")
plt.xlabel("Semana")
plt.ylabel("Avance acumulado (%)")
plt.legend()
plt.grid(True, alpha=0.3)
plt.tight_layout()
plt.savefig(ruta_curva_s, dpi=120)
plt.close()


# Crear y guardar el gráfico de producción semanal
plt.figure(figsize=(8, 4.5))
plt.plot(
    avance_semanal["semana"],
    avance_semanal["avance_plan_semanal"],
    marker="o",
    label="Planificado"
)
plt.plot(
    avance_semanal["semana"],
    avance_semanal["avance_real_semanal"],
    marker="s",
    label="Real"
)
plt.title("Producción semanal - Planificado vs. real")
plt.xlabel("Semana")
plt.ylabel("Avance semanal (%)")
plt.legend()
plt.grid(True, alpha=0.3)
plt.tight_layout()
plt.savefig(ruta_produccion, dpi=120)
plt.close()


# Crear y guardar el gráfico de avance real por especialidad
plt.figure(figsize=(10, 5))
barras = plt.bar(
    avance_especialidad["especialidad"],
    avance_especialidad["avance_real_semanal"]
)
plt.title("Avance real por especialidad")
plt.xlabel("Especialidad")
plt.ylabel("Avance real (%)")
plt.xticks(rotation=20)
plt.grid(True, axis="y", alpha=0.3)

# Escribir los valores encima de las barras
for barra, valor in zip(
    barras,
    avance_especialidad["avance_real_semanal"]
):
    plt.text(
        barra.get_x() + barra.get_width() / 2,
        valor + 0.2,
        f"{valor:.2f}",
        ha="center"
    )

# Ajustar y guardar
plt.tight_layout()
plt.savefig(ruta_especialidad, dpi=120)
plt.close()


# Crear y guardar el gráfico de cantidad de partidas por encargado
plt.figure(figsize=(9, 5))
barras = plt.bar(
    partidas_encargado["encargado"],
    partidas_encargado["cantidad_partidas"]
)
plt.title("Cantidad de partidas por encargado")
plt.xlabel("Encargado")
plt.ylabel("Cantidad de partidas")
plt.xticks(rotation=20)
plt.grid(True, axis="y", alpha=0.3)

# Escribir los valores encima de las barras
for barra, valor in zip(
    barras,
    partidas_encargado["cantidad_partidas"]
):
    plt.text(
        barra.get_x() + barra.get_width() / 2,
        valor + 0.05,
        str(valor),
        ha="center"
    )

# Ajustar y guardar
plt.tight_layout()
plt.savefig(ruta_partidas_encargado, dpi=120)
plt.close()


# Definir posiciones para el gráfico de avance por encargado
x = np.arange(len(avance_encargado["encargado"]))

# Definir el ancho de las barras
ancho = 0.35

# Crear las barras planificadas y reales
plt.figure(figsize=(10, 5))
plt.bar(
    x - ancho / 2,
    avance_encargado["avance_plan_semanal"],
    width=ancho,
    label="Planificado"
)
plt.bar(
    x + ancho / 2,
    avance_encargado["avance_real_semanal"],
    width=ancho,
    label="Real"
)
plt.xticks(
    x,
    avance_encargado["encargado"],
    rotation=20
)
plt.title("Avance planificado vs. real por encargado")
plt.xlabel("Encargado")
plt.ylabel("Avance (%)")
plt.legend()
plt.grid(True, axis="y", alpha=0.3)
plt.tight_layout()
plt.savefig(ruta_avance_encargado, dpi=120)
plt.close()

# Confirmar la generación de los gráficos
print("\nGráficos estáticos generados correctamente.")


# ============================================================
# 8. REPORTE EXCEL CON OPENPYXL
# ============================================================

# Exportar los resultados a diferentes hojas de Excel
with pd.ExcelWriter(ruta_excel, engine="openpyxl") as writer:
    # Exportar la base original
    df.to_excel(
        writer,
        sheet_name="Base_Datos",
        index=False
    )

    # Exportar el resumen semanal
    resumen_semanal.to_excel(
        writer,
        sheet_name="Resumen_Semanal",
        index=False
    )

    # Exportar el análisis por especialidad
    avance_especialidad.to_excel(
        writer,
        sheet_name="Especialidades",
        index=False
    )

    # Exportar las partidas por encargado
    partidas_encargado.to_excel(
        writer,
        sheet_name="Partidas_Encargado",
        index=False
    )

    # Exportar el avance por encargado
    avance_encargado.to_excel(
        writer,
        sheet_name="Avance_Encargado",
        index=False
    )


# Abrir el archivo Excel para aplicar formato
wb = load_workbook(ruta_excel)

# Definir estilos para los encabezados
relleno_encabezado = PatternFill("solid", fgColor="14253C")
fuente_encabezado = Font(bold=True, color="FFFFFF")
alineacion_centro = Alignment(horizontal="center", vertical="center")

# Definir bordes
lado = Side(style="thin", color="D9D9D9")
borde = Border(left=lado, right=lado, top=lado, bottom=lado)

# Recorrer las hojas y aplicar formato
for hoja in wb.worksheets:
    # Dar formato a la primera fila
    for celda in hoja[1]:
        celda.font = fuente_encabezado
        celda.fill = relleno_encabezado
        celda.alignment = alineacion_centro

    # Ajustar los anchos y aplicar bordes
    for columna in hoja.columns:
        max_longitud = 0

        for celda in columna:
            celda.border = borde

            if celda.value is not None:
                max_longitud = max(
                    max_longitud,
                    len(str(celda.value))
                )

        # Limitar el ancho máximo para mantener el reporte ordenado
        ancho_columna = min(max_longitud + 2, 35)

        # Aplicar el ancho
        hoja.column_dimensions[
            columna[0].column_letter
        ].width = ancho_columna


# Crear una hoja Dashboard
if "Dashboard" in wb.sheetnames:
    # Eliminar el Dashboard anterior si ya existía
    del wb["Dashboard"]

# Crear una nueva hoja Dashboard
hoja_dashboard = wb.create_sheet("Dashboard")

# Escribir el título principal
hoja_dashboard["A1"] = "REPORTE DE AVANCE FISICO DE OBRA"

# Dar formato al título
hoja_dashboard["A1"].font = Font(
    bold=True,
    size=14,
    color="14253C"
)

# Insertar la Curva S
imagen_curva = Image(ruta_curva_s)
imagen_curva.anchor = "A3"
hoja_dashboard.add_image(imagen_curva)

# Insertar el gráfico de producción semanal
imagen_produccion = Image(ruta_produccion)
imagen_produccion.anchor = "A25"
hoja_dashboard.add_image(imagen_produccion)

# Guardar el reporte final
wb.save(ruta_excel)

# Mostrar confirmación
print("Reporte Excel generado correctamente:")
print(ruta_excel)


# ============================================================
# 9. INTERPRETACION FINAL
# ============================================================

# Obtener los resultados de la última semana
resultado_final = resumen_semanal.iloc[-1]

# Mostrar los indicadores finales
print("\n===== RESULTADOS FINALES =====")
print(
    f"Avance planificado final: "
    f"{resultado_final['avance_plan_acumulado']:.2f} %"
)
print(
    f"Avance real final: "
    f"{resultado_final['avance_real_acumulado']:.2f} %"
)
print(
    f"Diferencia final: "
    f"{resultado_final['diferencia']:.2f} puntos porcentuales"
)
print(
    f"Estado general de la obra: "
    f"{resultado_final['estado']}"
)

# Ordenar las especialidades por diferencia
especialidades_ordenadas = avance_especialidad.sort_values("diferencia")

# Mostrar la especialidad con mayor desviación
print("\nEspecialidad con mayor desviación:")
print(especialidades_ordenadas.iloc[0])

# Ordenar los encargados por diferencia
encargados_ordenados = avance_encargado.sort_values("diferencia")

# Mostrar el encargado con mayor desviación
print("\nEncargado con mayor desviación:")
print(encargados_ordenados.iloc[0])


# ============================================================
# 10. GRAFICO ANIMADO DE LA CURVA S
# ============================================================

# Crear la figura para la animación
fig_anim, ax_anim = plt.subplots(figsize=(9, 5))

# Crear la línea planificada
linea_plan, = ax_anim.plot(
    [],
    [],
    marker="o",
    linestyle="--",
    linewidth=3,
    label="Planificado"
)

# Crear la línea real
linea_real, = ax_anim.plot(
    [],
    [],
    marker="s",
    linestyle="-",
    linewidth=2,
    label="Real"
)

# Configurar los límites del gráfico
ax_anim.set_xlim(1, 12)
ax_anim.set_ylim(0, 105)

# Agregar etiquetas
ax_anim.set_xlabel("Semana")
ax_anim.set_ylabel("Avance acumulado (%)")

# Agregar cuadrícula y leyenda
ax_anim.grid(True, alpha=0.3)
ax_anim.legend()


# Crear la función que actualizará la animación
def actualizar_animacion(frame):
    # Seleccionar los datos hasta la semana actual
    semanas = resumen_semanal["semana"].iloc[:frame + 1]
    plan = resumen_semanal["avance_plan_acumulado"].iloc[:frame + 1]
    real = resumen_semanal["avance_real_acumulado"].iloc[:frame + 1]

    # Actualizar las curvas
    linea_plan.set_data(semanas, plan)
    linea_real.set_data(semanas, real)

    # Actualizar el título
    ax_anim.set_title(
        f"Evolucion de la Curva S - Semana {frame + 1}"
    )

    # Devolver los elementos modificados
    return linea_plan, linea_real


# Crear la animación
animacion = FuncAnimation(
    fig_anim,
    actualizar_animacion,
    frames=len(resumen_semanal),
    interval=700,
    repeat=True
)

# Mostrar la animación
plt.show()


# ============================================================
# FIN DEL PROGRAMA
# ============================================================

# Mostrar un mensaje final
print("\nProceso finalizado correctamente.")
