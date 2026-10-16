"""Regresión respecto al ET1, entradas, Excel e interacciones de Streamlit."""

import ast
import hashlib
from io import BytesIO
from pathlib import Path
import sys
import unittest

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from pandas.testing import assert_frame_equal
from streamlit.testing.v1 import AppTest

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from modulos import calculos
from modulos.datos import agregar_registro, leer_csv, validar_datos
from modulos.exportacion import generar_excel
from modulos.graficos import generar_graficos

NOMBRES = [
    "calcular_avance_semanal", "calcular_avance_acumulado", "calcular_diferencia_avance",
    "analizar_avance_especialidad", "contar_partidas_encargado", "analizar_avance_encargado",
]
FUENTE = RAIZ / "referencias/entregable_1/analisis_avance_fisico_obra.py"


class PruebasAnalisis(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.datos = leer_csv(RAIZ / "datos/base_datos_avance_obra.csv")

    def test_archivos_originales_intactos(self):
        self.assertEqual(hashlib.sha256(FUENTE.read_bytes()).hexdigest(),
                         "ebafb3523023e980daf4b782d12cf438d8ac1d2966cc835a9004f64305de31fe")
        self.assertEqual(hashlib.sha256((RAIZ / "datos/base_datos_avance_obra.csv").read_bytes()).hexdigest(),
                         "92e6588594c52b7e9b46d7c64f3f3bc4d9335a67ec69461fe650f1e51834bc6d")

    def test_seis_funciones_originales_y_todas_sus_tablas(self):
        original = ast.parse(FUENTE.read_text(encoding="utf-8-sig"))
        nuevo = ast.parse((RAIZ / "modulos/calculos.py").read_text(encoding="utf-8"))
        funciones = [n for n in original.body if isinstance(n, ast.FunctionDef) and n.name in NOMBRES]
        actuales = {n.name: n for n in nuevo.body if isinstance(n, ast.FunctionDef)}
        for funcion in funciones:
            with self.subTest(funcion=funcion.name):
                self.assertEqual(ast.dump(funcion), ast.dump(actuales[funcion.name]))
        # Ejecutar solo funciones, sin ejecutar las escrituras y ventanas del ET1.
        espacio = {"pd": pd, "np": np}
        exec(compile(ast.Module(body=funciones, type_ignores=[]), str(FUENTE), "exec"), espacio)
        resultados = calculos.analizar_obra(self.datos)
        semanal = espacio["calcular_diferencia_avance"](
            espacio["calcular_avance_acumulado"](espacio["calcular_avance_semanal"](self.datos))
        )
        assert_frame_equal(resultados["semanal"], semanal)
        for tabla, funcion in [
            ("especialidades", "analizar_avance_especialidad"),
            ("encargados", "analizar_avance_encargado"),
            ("partidas_encargado", "contar_partidas_encargado"),
        ]:
            assert_frame_equal(resultados[tabla], espacio[funcion](self.datos))

    def test_indicadores_base(self):
        k = calculos.analizar_obra(self.datos)["indicadores"]
        self.assertEqual((k["partidas"], k["semanas"]), (20, 12))
        self.assertAlmostEqual(k["plan"], 99.99)
        self.assertAlmostEqual(k["real"], 98.83)
        self.assertAlmostEqual(k["diferencia"], -1.16)
        self.assertEqual(k["estado"], "RETRASADA")

    def test_corte_temporal_y_sin_datos(self):
        r = calculos.analizar_obra(self.datos, 6)
        self.assertEqual(r["indicadores"]["semana"], 6)
        self.assertEqual(len(r["datos"]), 120)
        self.assertEqual(len(r["semanal"]), 6)
        with self.assertRaisesRegex(ValueError, "No hay registros"):
            calculos.analizar_obra(self.datos, 0)

    def test_csv_bytes_y_plantilla(self):
        datos = leer_csv((RAIZ / "datos/base_datos_avance_obra.csv").read_bytes())
        assert_frame_equal(datos, self.datos)
        parcial = leer_csv(RAIZ / "datos/plantilla_carga.csv")
        self.assertEqual(len(parcial), 2)
        self.assertTrue(calculos.analizar_obra(parcial)["alertas"])
        with self.assertRaises(ValueError):
            leer_csv(b"")

    def test_bases_de_prueba_tienen_equipos_distintos(self):
        conocidos = set(self.datos["encargado"])
        bases = sorted((RAIZ / "datos").glob("0[1-5]_*.csv"))
        self.assertEqual(len(bases), 5)
        for archivo, cantidad in zip(bases, [10, 7, 8, 5, 5]):
            with self.subTest(archivo=archivo.name):
                datos = leer_csv(archivo)
                encargados = set(datos["encargado"])
                self.assertEqual(len(encargados), cantidad)
                self.assertFalse(encargados & conocidos)
                conocidos.update(encargados)

    def test_filtros_combinados_conservan_aportes_y_base(self):
        original = self.datos.copy(deep=True)
        fila = self.datos.iloc[0]
        r = calculos.analizar_obra(
            self.datos, 6, especialidades=[fila["especialidad"]],
            encargados=[fila["encargado"]], partidas=[fila["codigo_partida"]],
        )
        esperados = self.datos.loc[
            (self.datos["codigo_partida"] == fila["codigo_partida"]) & (self.datos["semana"] <= 6)
        ].reset_index(drop=True)
        assert_frame_equal(r["datos"].reset_index(drop=True), esperados)
        semanal = calculos.calcular_diferencia_avance(calculos.calcular_avance_acumulado(
            calculos.calcular_avance_semanal(esperados)
        ))
        assert_frame_equal(r["semanal"], semanal)
        self.assertAlmostEqual(r["indicadores"]["real"], semanal.iloc[-1]["avance_real_acumulado"])
        self.assertAlmostEqual(r["parametros"]["peso_seleccion"], fila["peso_partida"])
        self.assertTrue(r["parametros"]["filtros_activos"])
        self.assertEqual(r["alertas"], [])
        self.assertEqual(list(generar_graficos(r)["curva_s"].data[1].y),
                         semanal["avance_real_acumulado"].tolist())
        assert_frame_equal(self.datos, original)

    def test_filtro_sin_coincidencias(self):
        with self.assertRaisesRegex(ValueError, "No hay registros"):
            calculos.analizar_obra(self.datos, 6, especialidades=["Especialidad inexistente"])

    def test_excel_registra_filtros_y_exporta_solo_seleccion(self):
        fila = self.datos.iloc[0]
        r = calculos.analizar_obra(self.datos, 6, partidas=[fila["codigo_partida"]])
        libro = load_workbook(BytesIO(generar_excel(r)))
        hoja = libro["Base_Datos"]
        self.assertEqual(hoja.max_row, len(r["datos"]) + 1)
        self.assertEqual({hoja.cell(f, 2).value for f in range(2, hoja.max_row + 1)},
                         {fila["codigo_partida"]})
        self.assertEqual(libro["Dashboard"]["D4"].value, 6)
        self.assertEqual(libro["Dashboard"]["A14"].value, f"Partidas: {fila['codigo_partida']}")
        self.assertAlmostEqual(libro["Dashboard"]["D6"].value, r["indicadores"]["real"])

    def test_validaciones_de_entrada(self):
        casos = [
            ("semana", np.inf), ("semana", 0), ("semana", 1.5),
            ("avance_real_semanal", -1), ("avance_plan_semanal", "texto"),
            ("peso_partida", 0), ("peso_partida", 101), ("encargado", " "),
        ]
        for columna, valor in casos:
            with self.subTest(columna=columna, valor=valor):
                datos = self.datos.copy().astype({columna: object})
                datos.loc[0, columna] = valor
                with self.assertRaises(ValueError):
                    validar_datos(datos)
        for datos in [self.datos.iloc[:0], self.datos.drop(columns="semana")]:
            with self.assertRaises(ValueError):
                validar_datos(datos)

    def test_duplicados_metadatos_y_excesos(self):
        with self.assertRaisesRegex(ValueError, "duplicados"):
            validar_datos(pd.concat([self.datos, self.datos.iloc[:1]]))
        metadatos = self.datos.copy()
        metadatos.loc[0, "encargado"] = "Otro responsable"
        with self.assertRaisesRegex(ValueError, "constante"):
            validar_datos(metadatos)
        exceso = self.datos.copy()
        exceso.loc[0, "avance_real_semanal"] = 200
        with self.assertRaisesRegex(ValueError, "supera"):
            validar_datos(exceso)

    def test_ingreso_manual_sin_modificar_original(self):
        original = self.datos.copy(deep=True)
        registro = self.datos.iloc[0].to_dict()
        registro.update(semana=13, avance_plan_semanal=0.0, avance_real_semanal=0.0)
        actualizado = agregar_registro(self.datos, registro)
        self.assertEqual(len(actualizado), 241)
        assert_frame_equal(self.datos, original)

    def test_excel_con_tablas_y_cinco_graficos(self):
        r = calculos.analizar_obra(self.datos, 6)
        libro = load_workbook(BytesIO(generar_excel(r)))
        self.assertEqual(libro.sheetnames, ["Dashboard", "Base_Datos", "Resumen_Semanal",
                                           "Especialidades", "Partidas_Encargado", "Avance_Encargado"])
        self.assertEqual(libro["Base_Datos"].max_row, 121)
        self.assertEqual(len(libro["Dashboard"]._charts), 5)
        self.assertAlmostEqual(libro["Dashboard"]["D6"].value, r["indicadores"]["real"])
        semanal = libro["Resumen_Semanal"]
        self.assertAlmostEqual(semanal.cell(semanal.max_row, 5).value, r["indicadores"]["real"])
        self.assertEqual(len(generar_graficos(r)), 5)
        # Un nombre ingresado como '=...' debe conservarse como texto, no fórmula.
        r["datos"] = r["datos"].copy()
        r["datos"].loc[r["datos"].index[0], "partida"] = "=1+1"
        libro = load_workbook(BytesIO(generar_excel(r)))
        self.assertEqual(libro["Base_Datos"]["C2"].data_type, "s")


class PruebasInterfaz(unittest.TestCase):
    def abrir(self):
        return AppTest.from_file(str(RAIZ / "app.py"), default_timeout=30).run()

    def test_inicio_y_semana_de_corte(self):
        app = self.abrir()
        self.assertFalse(app.exception)
        self.assertEqual([m.value for m in app.metric[:3]], ["99.99 %", "98.83 %", "-1.16 pp"])
        self.assertEqual(len(app.number_input) + len(app.text_input), 8)
        self.assertEqual(len(app.get("plotly_chart")), 5)
        app.selectbox(key="semana_corte").set_value(6).run()
        self.assertFalse(app.exception)
        datos = leer_csv(RAIZ / "datos/base_datos_avance_obra.csv")
        esperado = calculos.analizar_obra(datos, 6)["indicadores"]["real"]
        self.assertEqual(app.metric[1].value, f"{esperado:.2f} %")

    def test_formulario_valido_y_restauracion(self):
        app = self.abrir()
        fila = app.session_state["datos"].iloc[0]
        app.number_input(key="manual_semana").set_value(13)
        for clave, columna in [("manual_codigo", "codigo_partida"), ("manual_partida", "partida"),
                               ("manual_especialidad", "especialidad"), ("manual_encargado", "encargado")]:
            app.text_input(key=clave).set_value(str(fila[columna]))
        app.number_input(key="manual_peso").set_value(float(fila["peso_partida"]))
        next(b for b in app.button if b.label == "Agregar registro").click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        self.assertEqual(len(app.session_state["datos"]), 241)
        self.assertEqual(app.selectbox(key="semana_corte").value, 13)
        app.button(key="restaurar").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.session_state["datos"]), 240)
        self.assertEqual(app.selectbox(key="semana_corte").value, 12)

    def test_formulario_invalido_conserva_base(self):
        app = self.abrir()
        next(b for b in app.button if b.label == "Agregar registro").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(app.error)
        self.assertEqual(len(app.session_state["datos"]), 240)

    def test_cambiar_csv_detecta_equipo_y_reinicia_filtros(self):
        app = self.abrir()
        for archivo in sorted((RAIZ / "datos").glob("0[1-5]_*.csv")):
            with self.subTest(archivo=archivo.name):
                anteriores = app.session_state["datos"].copy()
                fila = anteriores.iloc[0]
                app.multiselect(key="analisis_especialidades").set_value([fila["especialidad"]]).run()
                app.multiselect(key="analisis_encargados").set_value([fila["encargado"]]).run()
                app.multiselect(key="analisis_partidas").set_value([fila["codigo_partida"]]).run()
                app.selectbox(key="filtro_encargado").set_value(fila["encargado"]).run()
                app.selectbox(key="semana_corte").set_value(1).run()
                app.file_uploader(key="archivo_csv").set_value(
                    (archivo.name, archivo.read_bytes(), "text/csv")
                ).run()
                # Seleccionar el archivo lo prepara; el botón confirma su carga.
                assert_frame_equal(app.session_state["datos"], anteriores)
                app.button(key="cargar_archivo").click().run()
                self.assertFalse(app.exception)
                self.assertFalse(app.error)
                datos = leer_csv(archivo)
                nombres = sorted(datos["encargado"].unique().tolist())
                assert_frame_equal(app.session_state["datos"], datos)
                self.assertEqual(app.multiselect(key="analisis_encargados").options, nombres)
                self.assertEqual(app.selectbox(key="filtro_encargado").options, ["Todos"] + nombres)
                self.assertEqual(app.selectbox(key="filtro_encargado").value, "Todos")
                self.assertTrue(all(not widget.value for widget in app.multiselect))
                self.assertEqual(app.selectbox(key="semana_corte").value, int(datos["semana"].max()))
                self.assertEqual(app.session_state["fuente"], f"Archivo: {archivo.name}")
                self.assertTrue(any(f"{len(nombres)} encargados detectados" in c.value for c in app.caption))
                esperado = calculos.analizar_obra(datos)["indicadores"]["real"]
                self.assertEqual(app.metric[1].value, f"{esperado:.2f} %")
                tablas = [d.value for d in app.dataframe if "encargado" in d.value.columns]
                self.assertTrue(tablas)
                for tabla in tablas:
                    self.assertEqual(set(tabla["encargado"]), set(nombres))
        app.button(key="restaurar").click().run()
        self.assertFalse(app.exception)
        original = leer_csv(RAIZ / "datos/base_datos_avance_obra.csv")
        self.assertEqual(app.multiselect(key="analisis_encargados").options,
                         sorted(original["encargado"].unique().tolist()))

    def test_carga_csv_invalida_conserva_datos_y_filtros(self):
        app = self.abrir()
        datos = app.session_state["datos"].copy()
        nombre = datos.iloc[0]["encargado"]
        app.multiselect(key="analisis_encargados").set_value([nombre]).run()
        app.file_uploader(key="archivo_csv").set_value(
            ("invalido.csv", b"columna_desconocida\n1\n", "text/csv")
        ).run()
        app.button(key="cargar_archivo").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(app.error)
        self.assertEqual(app.multiselect(key="analisis_encargados").value, [nombre])
        assert_frame_equal(app.session_state["datos"], datos)

    def test_filtros_de_interfaz_y_limpieza(self):
        app = self.abrir()
        datos = app.session_state["datos"].copy()
        fila = datos.iloc[0]
        app.multiselect(key="analisis_especialidades").set_value([fila["especialidad"]]).run()
        app.multiselect(key="analisis_encargados").set_value([fila["encargado"]]).run()
        app.multiselect(key="analisis_partidas").set_value([fila["codigo_partida"]]).run()
        app.selectbox(key="semana_corte").set_value(6).run()
        self.assertFalse(app.exception)
        esperado = calculos.analizar_obra(datos, 6, partidas=[fila["codigo_partida"]])
        self.assertEqual(app.metric[1].value, f"{esperado['indicadores']['real']:.2f} %")
        self.assertEqual(len(app.get("plotly_chart")), 5)
        assert_frame_equal(app.session_state["datos"], datos)
        app.button(key="limpiar_filtros").click().run()
        self.assertFalse(app.exception)
        self.assertEqual([m.value for m in app.metric[:3]], ["99.99 %", "98.83 %", "-1.16 pp"])
        self.assertEqual(app.selectbox(key="semana_corte").value, 12)
        self.assertTrue(all(not widget.value for widget in app.multiselect))

    def test_filtros_dependientes_eliminan_seleccion_incompatible(self):
        app = self.abrir()
        datos = app.session_state["datos"]
        fila = datos.iloc[0]
        otra = next(e for e in datos["especialidad"].unique() if e != fila["especialidad"])
        app.multiselect(key="analisis_partidas").set_value([fila["codigo_partida"]]).run()
        app.multiselect(key="analisis_especialidades").set_value([otra]).run()
        self.assertFalse(app.exception)
        self.assertEqual(app.multiselect(key="analisis_partidas").value, [])
        esperado = calculos.analizar_obra(datos, 12, especialidades=[otra])
        self.assertEqual(app.metric[1].value, f"{esperado['indicadores']['real']:.2f} %")

    def test_seleccion_sin_registros_se_puede_restaurar(self):
        app = self.abrir()
        datos = app.session_state["datos"].iloc[:2].copy()
        datos.loc[datos.index[1], "semana"] = 2
        datos.loc[datos.index[1], "especialidad"] = "Solo semana 2"
        app.session_state["datos"] = validar_datos(datos)
        app.session_state["ajustar_corte"] = True
        app.run()
        app.multiselect(key="analisis_especialidades").set_value(["Solo semana 2"]).run()
        app.selectbox(key="semana_corte").set_value(1).run()
        self.assertFalse(app.exception)
        self.assertTrue(any("No hay registros" in aviso.value for aviso in app.info))
        app.button(key="limpiar_filtros").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.metric), 4)
        self.assertEqual(app.selectbox(key="semana_corte").value, 2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
