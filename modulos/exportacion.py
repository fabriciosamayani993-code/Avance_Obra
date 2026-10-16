"""Reporte Excel en memoria: conserva las tablas del primer entregable."""

from io import BytesIO

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


def generar_excel(resultados):
    """Devolver un XLSX con seis hojas, indicadores y cinco gráficos nativos.

    Los resultados los calcula Python. Excel permite consultar las tablas y
    gráficos; para recalcular con nuevos datos se utiliza la aplicación web.
    """
    memoria = BytesIO()
    hojas = {
        "Base_Datos": resultados["datos"],
        "Resumen_Semanal": resultados["semanal"],
        "Especialidades": resultados["especialidades"],
        "Partidas_Encargado": resultados["partidas_encargado"],
        "Avance_Encargado": resultados["encargados"],
    }
    with pd.ExcelWriter(memoria, engine="openpyxl") as escritor:
        for nombre, tabla in hojas.items():
            tabla.to_excel(escritor, sheet_name=nombre, index=False)
        libro = escritor.book
        for hoja in libro.worksheets:
            hoja.freeze_panes = "C2" if hoja.title == "Base_Datos" else "B2"
            hoja.auto_filter.ref = hoja.dimensions
            hoja.row_dimensions[1].height = 44
            for celda in hoja[1]:
                celda.fill = PatternFill("solid", fgColor="18344B")
                celda.font = Font(name="Calibri", bold=True, color="FFFFFF")
                celda.alignment = Alignment(wrap_text=True, vertical="center")
            for columna in hoja.columns:
                letra = columna[0].column_letter
                ancho = max(len(str(c.value or "")) for c in columna[1:]) + 3
                hoja.column_dimensions[letra].width = min(max(ancho, 20), 44)
                for celda in columna[1:]:
                    celda.font = Font(name="Calibri", size=11, color="18344B")
                    if celda.row % 2 == 0:
                        celda.fill = PatternFill("solid", fgColor="F0F5FA")
                    if isinstance(celda.value, (int, float)):
                        celda.number_format = (
                            "0" if columna[0].value in ["semana", "cantidad_partidas"] else "0.00"
                        )
                    # Los nombres y códigos ingresados se guardan como texto.
                    if isinstance(celda.value, str):
                        celda.data_type = "s"
        base = libro["Base_Datos"]
        base.column_dimensions["C"].width = 44
        dashboard = libro.create_sheet("Dashboard", 0)
        dashboard.sheet_view.showGridLines = False
        dashboard.merge_cells("A1:L2")
        dashboard["A1"] = "CONTROL DEL AVANCE FÍSICO DE OBRA"
        dashboard["A1"].font = Font(name="Calibri", size=20, bold=True, color="FFFFFF")
        dashboard["A1"].fill = PatternFill("solid", fgColor="18344B")
        dashboard["A1"].alignment = Alignment(vertical="center")
        indicadores = resultados["indicadores"]
        for fila, (etiqueta, valor, formato) in enumerate([
            ("Semana de corte", indicadores["semana"], "0"),
            ("Avance planificado (%)", indicadores["plan"], '0.00" %"'),
            ("Avance real (%)", indicadores["real"], '0.00" %"'),
            ("Desviación (real − plan)", indicadores["diferencia"], '0.00" pp"'),
            ("Estado", indicadores["estado"], "General"),
        ], start=4):
            dashboard.cell(fila, 1, etiqueta).font = Font(bold=True, color="18344B")
            celda = dashboard.cell(fila, 4, valor)
            celda.font = Font(bold=True, size=13, color="0D9488")
            celda.number_format = formato
        dashboard.merge_cells("A10:L10")
        dashboard["A10"] = "Avances ponderados respecto al total de la obra. Desviación en puntos porcentuales (pp)."
        dashboard["A10"].font = Font(italic=True, color="526579")
        dashboard.merge_cells("A11:L11")
        dashboard["A11"] = "Reporte de consulta. Modifica los datos y recalcula desde la aplicación Streamlit."
        dashboard["A11"].font = Font(italic=True, color="526579")
        parametros = resultados["parametros"]
        for fila, nombre, etiqueta in [
            (12, "especialidades", "Especialidades"),
            (13, "encargados", "Encargados"),
            (14, "partidas", "Partidas"),
        ]:
            dashboard.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=12)
            seleccion = ", ".join(parametros[nombre]) or "Todas las opciones disponibles"
            celda = dashboard.cell(fila, 1, f"{etiqueta}: {seleccion}")
            celda.font = Font(color="18344B")
            celda.alignment = Alignment(wrap_text=True, vertical="center")
            dashboard.row_dimensions[fila].height = 32
        dashboard.merge_cells("A15:L15")
        dashboard["A15"] = (f"Peso de la selección: {parametros['peso_seleccion']:.2f} %. "
                            "Los avances son aportes al total de la obra.")
        dashboard["A15"].font = Font(italic=True, color="526579")
        for letra in "ABCDEFGHIJKL":
            dashboard.column_dimensions[letra].width = 12
        dashboard.column_dimensions["A"].width = 16
        dashboard.column_dimensions["D"].width = 24
        for fila in range(4, 9):
            dashboard.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=3)
            dashboard.row_dimensions[fila].height = 23

        # Los gráficos se vinculan a las tablas exportadas, sin capturas de pantalla.
        configuraciones = [
            ("Resumen_Semanal", [4, 5], "Curva S", "A18", True, "Avance acumulado (%)"),
            ("Resumen_Semanal", [2, 3], "Producción semanal", "H18", False, "Aporte (pp)"),
            ("Especialidades", [2, 3], "Avance por especialidad", "A36", False, "Aporte (pp)"),
            ("Avance_Encargado", [2, 3], "Avance por encargado", "H36", False, "Aporte (pp)"),
            ("Partidas_Encargado", [2], "Partidas por encargado", "A54", False, "Partidas"),
        ]
        for origen, columnas, titulo, ancla, es_linea, unidad in configuraciones:
            hoja = libro[origen]
            grafico = LineChart() if es_linea else BarChart()
            grafico.title = titulo
            grafico.y_axis.title = unidad
            grafico.style = 13
            grafico.height, grafico.width = 8.3, 16.5
            for columna in columnas:
                grafico.add_data(Reference(hoja, min_col=columna, min_row=1, max_row=hoja.max_row),
                                 titles_from_data=True)
            grafico.set_categories(Reference(hoja, min_col=1, min_row=2, max_row=hoja.max_row))
            dashboard.add_chart(grafico, ancla)
        dashboard.freeze_panes = "A16"
    return memoria.getvalue()
